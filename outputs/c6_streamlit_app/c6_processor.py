from __future__ import annotations

import argparse
import csv
import json
import io
import re
import secrets
import sys
import time
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime
from http.server import BaseHTTPRequestHandler, HTTPServer
from pathlib import Path
from typing import BinaryIO, Iterable
from urllib import error as urllib_error
from urllib import parse as urllib_parse
from urllib import request as urllib_request
import webbrowser

import pandas as pd
import xlsxwriter
from xlsxwriter.utility import xl_rowcol_to_cell


COEF_VALUE = 0.0228
DEFAULT_PHONE_BASE_DIR = Path(
    r"G:\Meu Drive\GRUPO IMPÉRIO\LEADS\TELEFONES\HUB CONSOLIDADO"
)
FALLBACK_PHONE_BASE_DIR = Path(
    r"H:\Meu Drive\GRUPO IMPÉRIO\LEADS\TELEFONES\HUB CONSOLIDADO"
)
DEFAULT_PHONE_BASE_DIRS = (DEFAULT_PHONE_BASE_DIR, FALLBACK_PHONE_BASE_DIR)
ARGUS_DISC_BASE_DIRS = (
    Path(r"G:\Meu Drive\GRUPO IMPÉRIO\LEADS\ARGUS\DISC"),
    Path(r"H:\Meu Drive\GRUPO IMPÉRIO\LEADS\ARGUS\DISC"),
)
DDB_BASE_DIRS = (
    Path("G:\\Meu Drive\\GRUPO IMP\u00c9RIO\\LEADS\\INSS DT DDB"),
    Path("H:\\Meu Drive\\GRUPO IMP\u00c9RIO\\LEADS\\INSS DT DDB"),
)
DISPARO_PRO_ENVIADOS_DIR = Path(
    "G:\\Meu Drive\\GRUPO IMP\u00c9RIO\\LEADS\\DISPARO PRO\\Enviados"
)
DISPARO_PRO_AUTO_ENVIADOS_DIR = Path(
    "G:\\Meu Drive\\GRUPO IMP\u00c9RIO\\LEADS\\DISPARO PRO\\auto-enviados"
)
GOOGLE_DRIVE_AUTO_ENVIADOS_PATH = ("GRUPO IMP\u00c9RIO", "LEADS", "DISPARO PRO", "auto-enviados")
GOOGLE_DRIVE_AUTO_ENVIADOS_FOLDER_ID = "1ZBgiFOqO18Eez_LacdDAD8skvjD3tE_4"
GOOGLE_SCOPES = (
    "https://www.googleapis.com/auth/drive",
    "https://www.googleapis.com/auth/spreadsheets",
)
BRAZIL_UFS = {
    "AC",
    "AL",
    "AM",
    "AP",
    "BA",
    "CE",
    "DF",
    "ES",
    "GO",
    "MA",
    "MG",
    "MS",
    "MT",
    "PA",
    "PB",
    "PE",
    "PI",
    "PR",
    "RJ",
    "RN",
    "RO",
    "RR",
    "RS",
    "SC",
    "SE",
    "SP",
    "TO",
}
MONTH_NAMES_PT_BR = {
    1: "JANEIRO",
    2: "FEVEREIRO",
    3: "MARÇO",
    4: "ABRIL",
    5: "MAIO",
    6: "JUNHO",
    7: "JULHO",
    8: "AGOSTO",
    9: "SETEMBRO",
    10: "OUTUBRO",
    11: "NOVEMBRO",
    12: "DEZEMBRO",
}

SOURCE_COLUMNS = {
    "ovc_cpf": "CPF",
    "cr_matricula": "BENEFÍCIO",
    "cr_parcelas_contrato": "FEZ EM",
    "cr_parcelas_em_aberto": "FALTAM",
    "cr_valor_parcela": "PARCELA",
    "cr_saldo_refin": "SALDO DEV",
    "cr_taxa": "TAXA",
    "dc_nome_1": "NOME",
    "dc_data_nascimento_1": "DATA DE NASC",
    "tf_ddd_1": "DDD C6",
    "tf_telefone_1": "TEL C6",
}

IGNORED_COLUMNS = {"cr_valor_contrato", "cr_observacao"}

FINAL_COLUMNS = [
    "CPF",
    "BENEFÍCIO",
    "FEZ EM",
    "FALTAM",
    "PARCELA",
    "SALDO DEV",
    "TAXA",
    "NOME",
    "DATA DE NASC",
    "DDD C6",
    "TEL C6",
    "TELEFONE C6",
    "TELEFONE BASE",
    "COEF",
    "BRUTO",
    "TROCO",
    "BDD",
]

ARGUS_DISCADORA_COLUMNS = [
    "CAMPANHA",
    "CPF / CNPJ",
    "Taxa",
    "FEZ EM",
    "FALTAM",
    "PARCELA",
    "TROCO",
    "Telefone",
    "Nome",
    "Data Nascimento",
    "Telefone",
    "DDB",
    "Beneficio",
]

SMS_COLUMNS = ["Telefone", "Nome", "N", "BDD"]
SMS_ENVIADOS_COLUMNS = [
    "cpf",
    "benefício",
    "fez em",
    "parcela",
    "taxa",
    "nome",
    " troco ",
    "TELEFONE BASE",
    "BDD",
]


@dataclass(frozen=True)
class FilterOptions:
    faltam: tuple[int, ...] = ()
    faltam_modo: str = "maior"
    taxa_min: float = 0.0
    troco_min: float = 0.0
    data_inicial: date | None = date(1954, 1, 1)
    data_final: date | None = date(1995, 12, 31)
    telefone_data_inicial: date | None = None


class GoogleSheetsSetupError(RuntimeError):
    pass


def read_c6_csv(source: str | Path | bytes | BinaryIO) -> pd.DataFrame:
    data = _read_bytes(source)
    last_error: Exception | None = None

    for encoding in ("utf-8-sig", "cp1252", "latin1"):
        try:
            text = data.decode(encoding)
            df = pd.read_csv(io.StringIO(text), sep=";", dtype=str)
            df.columns = [str(col).strip().lstrip("\ufeff") for col in df.columns]
            return df
        except Exception as exc:
            last_error = exc

    raise ValueError(f"Não consegui ler o CSV. Último erro: {last_error}")


def transform_c6_dataframe(raw_df: pd.DataFrame) -> pd.DataFrame:
    normalized_columns = {_normalize_header(column): column for column in raw_df.columns}
    benefit_source = "cr_matricula" if "cr_matricula" in raw_df.columns else normalized_columns.get("BENEFICIO")
    missing = [
        col for col in SOURCE_COLUMNS
        if col != "cr_matricula" and col not in raw_df.columns
    ]
    if benefit_source is None:
        missing.append("cr_matricula ou Beneficio")
    if missing:
        joined = ", ".join(missing)
        raise ValueError(f"Cabeçalhos obrigatórios ausentes: {joined}")

    output = pd.DataFrame()
    output["CPF"] = raw_df["ovc_cpf"].map(_digits_only)
    output["BENEFÍCIO"] = raw_df[benefit_source].map(_digits_only)
    output["FEZ EM"] = _parse_int_series(raw_df["cr_parcelas_contrato"])
    output["FALTAM"] = _parse_int_series(raw_df["cr_parcelas_em_aberto"])
    output["PARCELA"] = raw_df["cr_valor_parcela"].map(_parse_br_number)
    output["SALDO DEV"] = raw_df["cr_saldo_refin"].map(_parse_br_number)
    output["TAXA"] = raw_df["cr_taxa"].map(_parse_br_number)
    output["NOME"] = raw_df["dc_nome_1"].fillna("").astype(str).str.strip()
    output["DATA DE NASC"] = pd.to_datetime(
        raw_df["dc_data_nascimento_1"], dayfirst=True, errors="coerce"
    )
    output["DDD C6"] = raw_df["tf_ddd_1"].map(_digits_only)
    output["TEL C6"] = raw_df["tf_telefone_1"].map(_digits_only)
    output["TELEFONE C6"] = output["DDD C6"] + output["TEL C6"]
    output["TELEFONE BASE"] = ""
    output["COEF"] = COEF_VALUE
    output["BRUTO"] = output["PARCELA"] / output["COEF"]
    output["TROCO"] = output["BRUTO"] - output["SALDO DEV"]
    output["BDD"] = ""

    return output[FINAL_COLUMNS]


def apply_filters(df: pd.DataFrame, options: FilterOptions) -> pd.DataFrame:
    mask = df["SALDO DEV"].fillna(0) > 0

    if options.faltam:
        faltam_values = tuple(int(value) for value in options.faltam)
        if options.faltam_modo == "exato":
            mask &= df["FALTAM"].isin(faltam_values)
        elif options.faltam_modo == "ate":
            mask &= df["FALTAM"] <= max(faltam_values)
        elif options.faltam_modo == "a_partir":
            mask &= df["FALTAM"] >= min(faltam_values)
        else:
            mask &= df["FALTAM"] > min(faltam_values)

    if options.taxa_min is not None:
        mask &= df["TAXA"].fillna(float("-inf")) >= float(options.taxa_min)

    if options.troco_min is not None:
        mask &= df["TROCO"].fillna(float("-inf")) > float(options.troco_min)

    if options.data_inicial is not None:
        start = pd.Timestamp(options.data_inicial)
        mask &= df["DATA DE NASC"].notna() & (df["DATA DE NASC"] >= start)

    if options.data_final is not None:
        end = pd.Timestamp(options.data_final)
        mask &= df["DATA DE NASC"].notna() & (df["DATA DE NASC"] <= end)

    return df.loc[mask].copy().reset_index(drop=True)


def parse_faltam_values(text: str | Iterable[int] | None) -> tuple[int, ...]:
    if text is None:
        return ()

    if not isinstance(text, str):
        return tuple(int(value) for value in text)

    cleaned = text.strip()
    if not cleaned:
        return ()

    values: list[int] = []
    for part in re.split(r"[,\s;/]+", cleaned):
        if part:
            values.append(int(part))

    return tuple(dict.fromkeys(values))


def parse_date_value(value: str | date | None) -> date | None:
    if value is None or value == "":
        return None

    if isinstance(value, date):
        return value

    parsed = pd.to_datetime(str(value), dayfirst=True, errors="raise")
    return parsed.date()


def write_xlsx(df: pd.DataFrame, output_path: str | Path | BinaryIO) -> None:
    if hasattr(output_path, "write"):
        workbook = xlsxwriter.Workbook(output_path, {"in_memory": True})
    else:
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        workbook = xlsxwriter.Workbook(str(output_path), {"constant_memory": True})
    worksheet = workbook.add_worksheet("C6 Filtrado")

    header_fmt = workbook.add_format(
        {
            "bold": True,
            "font_color": "#FFFFFF",
            "bg_color": "#1F4E78",
            "border": 1,
            "border_color": "#B7C9D6",
            "align": "center",
            "valign": "vcenter",
        }
    )
    text_fmt = workbook.add_format({"num_format": "@", "valign": "vcenter"})
    digit_number_fmt = workbook.add_format({"num_format": "0", "valign": "vcenter"})
    int_fmt = workbook.add_format({"num_format": "#,##0", "valign": "vcenter"})
    number_fmt = workbook.add_format({"num_format": "0.00", "valign": "vcenter"})
    coef_fmt = workbook.add_format({"num_format": "0.0000", "valign": "vcenter"})
    money_fmt = workbook.add_format(
        {"num_format": '[$R$-pt-BR] #,##0.00', "valign": "vcenter"}
    )
    date_fmt = workbook.add_format({"num_format": "dd/mm/yyyy", "valign": "vcenter"})

    worksheet.set_tab_color("#1F4E78")
    worksheet.freeze_panes(1, 0)
    worksheet.set_row(0, 24)

    for col_idx, column in enumerate(FINAL_COLUMNS):
        worksheet.write(0, col_idx, column, header_fmt)

    col_index = {name: idx for idx, name in enumerate(FINAL_COLUMNS)}

    for row_idx, row in df.iterrows():
        excel_row = row_idx + 1
        _write_digits_number(worksheet, excel_row, col_index["CPF"], row["CPF"], digit_number_fmt)
        _write_digits_number(
            worksheet, excel_row, col_index["BENEFÍCIO"], row["BENEFÍCIO"], digit_number_fmt
        )
        _write_number(worksheet, excel_row, col_index["FEZ EM"], row["FEZ EM"], int_fmt)
        _write_number(worksheet, excel_row, col_index["FALTAM"], row["FALTAM"], int_fmt)
        _write_number(worksheet, excel_row, col_index["PARCELA"], row["PARCELA"], money_fmt)
        _write_number(
            worksheet, excel_row, col_index["SALDO DEV"], row["SALDO DEV"], money_fmt
        )
        _write_number(worksheet, excel_row, col_index["TAXA"], row["TAXA"], number_fmt)
        worksheet.write_string(excel_row, col_index["NOME"], _safe_text(row["NOME"]), text_fmt)
        _write_date(worksheet, excel_row, col_index["DATA DE NASC"], row["DATA DE NASC"], date_fmt)
        _write_digits_number(
            worksheet, excel_row, col_index["DDD C6"], row["DDD C6"], digit_number_fmt
        )
        _write_digits_number(
            worksheet, excel_row, col_index["TEL C6"], row["TEL C6"], digit_number_fmt
        )
        _write_digits_number(
            worksheet, excel_row, col_index["TELEFONE C6"], row["TELEFONE C6"], digit_number_fmt
        )
        telefone_base = _safe_text(row["TELEFONE BASE"]).strip()
        if telefone_base:
            worksheet.write_string(
                excel_row, col_index["TELEFONE BASE"], telefone_base, text_fmt
            )
        else:
            worksheet.write_blank(excel_row, col_index["TELEFONE BASE"], None, text_fmt)
        _write_number(worksheet, excel_row, col_index["COEF"], row["COEF"], coef_fmt)

        parcela_cell = xl_rowcol_to_cell(excel_row, col_index["PARCELA"])
        saldo_cell = xl_rowcol_to_cell(excel_row, col_index["SALDO DEV"])
        coef_cell = xl_rowcol_to_cell(excel_row, col_index["COEF"])
        bruto_cell = xl_rowcol_to_cell(excel_row, col_index["BRUTO"])
        worksheet.write_formula(
            excel_row,
            col_index["BRUTO"],
            f"={parcela_cell}/{coef_cell}",
            money_fmt,
            _number_or_blank(row["BRUTO"]),
        )
        worksheet.write_formula(
            excel_row,
            col_index["TROCO"],
            f"={bruto_cell}-{saldo_cell}",
            money_fmt,
            _number_or_blank(row["TROCO"]),
        )
        bdd = _format_br_date(row.get("BDD", ""))
        if bdd:
            worksheet.write_string(excel_row, col_index["BDD"], bdd, text_fmt)
        else:
            worksheet.write_blank(excel_row, col_index["BDD"], None, text_fmt)

    last_row = max(len(df), 1)
    worksheet.autofilter(0, 0, last_row, len(FINAL_COLUMNS) - 1)

    widths = {
        "CPF": 15,
        "BENEFÍCIO": 14,
        "FEZ EM": 10,
        "FALTAM": 10,
        "PARCELA": 13,
        "SALDO DEV": 14,
        "TAXA": 10,
        "NOME": 34,
        "DATA DE NASC": 14,
        "DDD C6": 9,
        "TEL C6": 13,
        "TELEFONE C6": 15,
        "TELEFONE BASE": 16,
        "COEF": 10,
        "BRUTO": 14,
        "TROCO": 14,
        "BDD": 12,
    }
    formats = {
        "CPF": digit_number_fmt,
        "BENEFÍCIO": digit_number_fmt,
        "FEZ EM": int_fmt,
        "FALTAM": int_fmt,
        "PARCELA": money_fmt,
        "SALDO DEV": money_fmt,
        "TAXA": number_fmt,
        "NOME": text_fmt,
        "DATA DE NASC": date_fmt,
        "BDD": text_fmt,
        "DDD C6": digit_number_fmt,
        "TEL C6": digit_number_fmt,
        "TELEFONE C6": digit_number_fmt,
        "TELEFONE BASE": text_fmt,
        "COEF": coef_fmt,
        "BRUTO": money_fmt,
        "TROCO": money_fmt,
    }

    for column, width in widths.items():
        idx = col_index[column]
        worksheet.set_column(idx, idx, width, formats[column])

    workbook.set_properties({"title": "C6 filtrado", "subject": "Exportação C6 tratada"})
    workbook.close()


def dataframe_to_xlsx_bytes(df: pd.DataFrame) -> bytes:
    buffer = io.BytesIO()
    write_xlsx(df, buffer)
    buffer.seek(0)
    return buffer.getvalue()


def build_sms_dataframe(df: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, str]] = []
    seen_phones: set[str] = set()
    for _, row in df.iterrows():
        nome = _safe_text(row.get("NOME", "")).strip()
        primeiro_nome = nome.split()[0] if nome.split() else ""
        for phone in _split_phone_values(row.get("TELEFONE BASE", "")):
            if phone in seen_phones:
                continue
            seen_phones.add(phone)
            rows.append(
                {
                    "Telefone": phone,
                    "Nome": nome,
                    "N": primeiro_nome,
                    "BDD": _format_br_date(row.get("BDD", "")),
                }
            )
    return pd.DataFrame(rows, columns=SMS_COLUMNS)


def write_sms_xlsx(df: pd.DataFrame, output_path: str | Path | BinaryIO) -> None:
    sms_df = build_sms_dataframe(df)
    if hasattr(output_path, "write"):
        workbook = xlsxwriter.Workbook(output_path, {"in_memory": True})
    else:
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        workbook = xlsxwriter.Workbook(str(output_path), {"constant_memory": True})

    worksheet = workbook.add_worksheet("SMS")
    header_fmt = workbook.add_format(
        {
            "bold": True,
            "font_color": "#FFFFFF",
            "bg_color": "#1F4E78",
            "border": 1,
            "border_color": "#B7C9D6",
            "align": "center",
            "valign": "vcenter",
        }
    )
    phone_fmt = workbook.add_format({"num_format": "00000000000", "valign": "vcenter"})
    text_fmt = workbook.add_format({"num_format": "@", "valign": "vcenter"})

    worksheet.set_tab_color("#2F75B5")
    worksheet.freeze_panes(1, 0)
    worksheet.set_row(0, 24)
    for col_idx, column in enumerate(SMS_COLUMNS):
        worksheet.write(0, col_idx, column, header_fmt)

    for row_idx, row in sms_df.iterrows():
        excel_row = row_idx + 1
        _write_digits_number(worksheet, excel_row, 0, row["Telefone"], phone_fmt)
        worksheet.write_string(excel_row, 1, _safe_text(row["Nome"]), text_fmt)
        worksheet.write_string(excel_row, 2, _safe_text(row["N"]), text_fmt)
        worksheet.write_string(excel_row, 3, _safe_text(row["BDD"]), text_fmt)

    last_row = max(len(sms_df), 1)
    worksheet.autofilter(0, 0, last_row, len(SMS_COLUMNS) - 1)
    worksheet.set_column(0, 0, 15, phone_fmt)
    worksheet.set_column(1, 1, 34, text_fmt)
    worksheet.set_column(2, 2, 18, text_fmt)
    worksheet.set_column(3, 3, 12, text_fmt)
    workbook.set_properties({"title": "SMS C6", "subject": "Base SMS C6"})
    workbook.close()


def sms_xlsx_bytes(df: pd.DataFrame) -> bytes:
    buffer = io.BytesIO()
    write_sms_xlsx(df, buffer)
    buffer.seek(0)
    return buffer.getvalue()


def build_sms_enviados_dataframe(df: pd.DataFrame) -> pd.DataFrame:
    output = pd.DataFrame(index=df.index)
    output["cpf"] = df["CPF"].map(_normalize_cpf_key)
    output["benefício"] = df[FINAL_COLUMNS[1]].map(_digits_only)
    output["fez em"] = df["FEZ EM"].map(_format_int)
    output["parcela"] = df["PARCELA"]
    output["taxa"] = df["TAXA"]
    output["nome"] = df["NOME"].map(lambda value: _safe_text(value).strip())
    output[" troco "] = df["TROCO"]
    output["TELEFONE BASE"] = df["TELEFONE BASE"].map(lambda value: _safe_text(value).strip())
    output["BDD"] = df["BDD"].map(_format_br_date) if "BDD" in df.columns else ""
    return output[SMS_ENVIADOS_COLUMNS]


def write_sms_enviados_xlsx(df: pd.DataFrame, output_path: str | Path | BinaryIO) -> None:
    enviados_df = build_sms_enviados_dataframe(df)
    if hasattr(output_path, "write"):
        workbook = xlsxwriter.Workbook(output_path, {"in_memory": True})
    else:
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        workbook = xlsxwriter.Workbook(str(output_path), {"constant_memory": True})

    worksheet = workbook.add_worksheet("enviados")
    header_fmt = workbook.add_format(
        {
            "bold": True,
            "font_color": "#FFFFFF",
            "bg_color": "#1F4E78",
            "border": 1,
            "border_color": "#B7C9D6",
            "align": "center",
            "valign": "vcenter",
        }
    )
    digit_fmt = workbook.add_format({"num_format": "0", "valign": "vcenter"})
    int_fmt = workbook.add_format({"num_format": "#,##0", "valign": "vcenter"})
    number_fmt = workbook.add_format({"num_format": "0.00", "valign": "vcenter"})
    money_fmt = workbook.add_format(
        {"num_format": '[$R$-pt-BR] #,##0.00', "valign": "vcenter"}
    )
    text_fmt = workbook.add_format({"num_format": "@", "valign": "vcenter"})

    worksheet.set_tab_color("#70AD47")
    worksheet.freeze_panes(1, 0)
    worksheet.set_row(0, 24)
    for col_idx, column in enumerate(SMS_ENVIADOS_COLUMNS):
        worksheet.write(0, col_idx, column, header_fmt)

    for row_idx, row in enviados_df.iterrows():
        excel_row = row_idx + 1
        _write_digits_number(worksheet, excel_row, 0, row["cpf"], digit_fmt)
        _write_digits_number(worksheet, excel_row, 1, row["benefício"], digit_fmt)
        _write_number(worksheet, excel_row, 2, row["fez em"], int_fmt)
        _write_number(worksheet, excel_row, 3, row["parcela"], money_fmt)
        _write_number(worksheet, excel_row, 4, row["taxa"], number_fmt)
        worksheet.write_string(excel_row, 5, _safe_text(row["nome"]), text_fmt)
        _write_number(worksheet, excel_row, 6, row[" troco "], money_fmt)
        worksheet.write_string(excel_row, 7, _safe_text(row["TELEFONE BASE"]), text_fmt)
        worksheet.write_string(excel_row, 8, _safe_text(row["BDD"]), text_fmt)

    last_row = max(len(enviados_df), 1)
    worksheet.autofilter(0, 0, last_row, len(SMS_ENVIADOS_COLUMNS) - 1)
    worksheet.set_column(0, 0, 15, digit_fmt)
    worksheet.set_column(1, 1, 14, digit_fmt)
    worksheet.set_column(2, 2, 10, int_fmt)
    worksheet.set_column(3, 3, 13, money_fmt)
    worksheet.set_column(4, 4, 10, number_fmt)
    worksheet.set_column(5, 5, 34, text_fmt)
    worksheet.set_column(6, 6, 14, money_fmt)
    worksheet.set_column(7, 7, 18, text_fmt)
    worksheet.set_column(8, 8, 12, text_fmt)
    workbook.set_properties({"title": "SMS enviados C6", "subject": "Base enviada SMS C6"})
    workbook.close()


def sms_enviados_xlsx_bytes(df: pd.DataFrame) -> bytes:
    buffer = io.BytesIO()
    write_sms_enviados_xlsx(df, buffer)
    buffer.seek(0)
    return buffer.getvalue()


def create_sms_enviados_google_sheet(df: pd.DataFrame, title: str) -> str:
    token = _load_google_access_token()
    folder_id = _google_auto_enviados_folder_id() or _ensure_drive_folder_path(
        token, GOOGLE_DRIVE_AUTO_ENVIADOS_PATH
    )

    file_metadata = {
        "name": title,
        "mimeType": "application/vnd.google-apps.spreadsheet",
        "parents": [folder_id],
    }
    created = _google_api_request(
        "POST",
        "https://www.googleapis.com/drive/v3/files?fields=id,webViewLink&supportsAllDrives=true",
        token,
        file_metadata,
    )
    spreadsheet_id = created["id"]

    _google_api_request(
        "POST",
        f"https://sheets.googleapis.com/v4/spreadsheets/{spreadsheet_id}:batchUpdate",
        token,
        {
            "requests": [
                {
                    "updateSheetProperties": {
                        "properties": {"sheetId": 0, "title": "enviados"},
                        "fields": "title",
                    }
                }
            ]
        },
    )

    enviados_df = build_sms_enviados_dataframe(df)
    values = [list(enviados_df.columns)]
    values.extend(
        [[_google_sheet_cell_value(value) for value in row] for row in enviados_df.itertuples(index=False)]
    )
    encoded_range = urllib_parse.quote("enviados!A1", safe="")
    _google_api_request(
        "PUT",
        f"https://sheets.googleapis.com/v4/spreadsheets/{spreadsheet_id}/values/{encoded_range}?valueInputOption=USER_ENTERED",
        token,
        {"values": values},
    )
    return created.get("webViewLink") or f"https://docs.google.com/spreadsheets/d/{spreadsheet_id}"


def build_sms_output_filename(
    source: str | Path | bytes | BinaryIO,
    row_count: int,
    base_dir: str | Path = DISPARO_PRO_ENVIADOS_DIR,
    today: date | None = None,
) -> str:
    campaign_number = find_next_sms_campaign_number(base_dir, today)
    list_name = build_sms_list_stem(source, row_count, campaign_number)
    return f"CAMP {campaign_number} {list_name}.xlsx"


def build_sms_output_path(
    source: str | Path | bytes | BinaryIO,
    row_count: int,
    base_dir: str | Path = DISPARO_PRO_ENVIADOS_DIR,
    today: date | None = None,
) -> Path:
    month_dir = resolve_sms_month_dir(base_dir, today)
    return month_dir / build_sms_output_filename(source, row_count, base_dir, today)


def build_sms_enviados_output_path(
    source: str | Path | bytes | BinaryIO,
    row_count: int,
    base_dir: str | Path = DISPARO_PRO_ENVIADOS_DIR,
    today: date | None = None,
) -> Path:
    auto_dir = resolve_sms_auto_enviados_dir()
    return auto_dir / build_sms_enviados_output_filename(source, row_count, base_dir, today)


def build_sms_enviados_output_filename(
    source: str | Path | bytes | BinaryIO,
    row_count: int,
    base_dir: str | Path = DISPARO_PRO_ENVIADOS_DIR,
    today: date | None = None,
) -> str:
    campaign_number = find_next_sms_campaign_number(base_dir, today)
    return f"{build_sms_list_stem(source, row_count, campaign_number)}.xlsx"


def build_sms_list_stem(
    source: str | Path | bytes | BinaryIO, row_count: int, campaign_number: int
) -> str:
    source_name = Path(str(getattr(source, "name", source))).stem
    uf = extract_uf_from_filename(source_name)
    return f"{row_count} refin c6 {uf} {campaign_number} (1 ou 0 cart\u00f5es) sem RL pra disc"


def build_argus_discadora_dataframe(df: pd.DataFrame, campanha: str = "") -> pd.DataFrame:
    required = ("CPF", "TAXA", "FEZ EM", "FALTAM", "PARCELA", "TROCO", "TELEFONE BASE",
                "NOME", "DATA DE NASC", "TELEFONE C6", "BDD", "BENEF\u00cdCIO")
    missing = [column for column in required if column not in df.columns]
    if missing:
        raise ValueError("Dados obrigatorios ausentes para o LAYOUT PADRAO C6: " + ", ".join(missing))

    values = [
        pd.Series(campanha, index=df.index),
        df["CPF"].map(_normalize_cpf_key),
        df["TAXA"].map(_format_br_decimal),
        df["FEZ EM"].map(_format_int),
        df["FALTAM"].map(_format_int),
        df["PARCELA"].map(_format_br_parcela),
        df["TROCO"].map(_format_br_accounting),
        df["TELEFONE BASE"].map(lambda value: _safe_text(value).strip()),
        df["NOME"].map(lambda value: _safe_text(value).strip()),
        df["DATA DE NASC"].map(_format_br_date),
        df["TELEFONE C6"].map(_format_identifier_exact),
        df["BDD"].map(_format_br_date),
        df["BENEF\u00cdCIO"].map(_digits_only),
    ]
    output = pd.concat(values, axis=1)
    output.columns = ARGUS_DISCADORA_COLUMNS
    return output


def argus_discadora_csv_bytes(df: pd.DataFrame, campanha: str = "") -> bytes:
    return _argus_dataframe_to_csv_bytes(build_argus_discadora_dataframe(df, campanha))


def write_argus_discadora_csv(
    df: pd.DataFrame, output_path: str | Path, campanha: str = ""
) -> None:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_bytes(argus_discadora_csv_bytes(df, campanha))


def build_argus_discadora_from_file(
    source: str | Path | bytes | BinaryIO, campanha: str = ""
) -> pd.DataFrame:
    if not campanha:
        campanha = build_argus_campaign_name(source)
    return build_argus_discadora_dataframe(read_processed_c6_file(source), campanha)


def write_argus_discadora_csv_from_file(
    source: str | Path | bytes | BinaryIO, output_path: str | Path, campanha: str = ""
) -> None:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_bytes(argus_discadora_csv_bytes_from_file(source, campanha))


def argus_discadora_csv_bytes_from_file(
    source: str | Path | bytes | BinaryIO, campanha: str = ""
) -> bytes:
    return _argus_dataframe_to_csv_bytes(build_argus_discadora_from_file(source, campanha))


def build_argus_campaign_name(
    source: str | Path | bytes | BinaryIO,
    base_dir: str | Path | Iterable[str | Path] = ARGUS_DISC_BASE_DIRS,
) -> str:
    source_name = Path(str(getattr(source, "name", source))).stem
    uf = extract_uf_from_filename(source_name)
    next_number = find_next_argus_campaign_number(uf, base_dir)
    return f"{uf} {next_number} REFIN INSS C6"


def build_argus_output_filename(
    source: str | Path | bytes | BinaryIO,
    base_dir: str | Path | Iterable[str | Path] = ARGUS_DISC_BASE_DIRS,
) -> str:
    source_stem = _clean_source_stem(Path(str(getattr(source, "name", source))).stem)
    return f"{build_argus_campaign_name(source, base_dir)} {source_stem}.csv"


def build_argus_output_path(
    source: str | Path | bytes | BinaryIO,
    base_dir: str | Path | Iterable[str | Path] = ARGUS_DISC_BASE_DIRS,
    today: date | None = None,
) -> Path:
    source_name = Path(str(getattr(source, "name", source))).stem
    uf = extract_uf_from_filename(source_name)
    resolved_base_dir = resolve_argus_disc_base_dir(base_dir)
    return (
        resolved_base_dir
        / _argus_month_folder_name(today)
        / uf
        / build_argus_output_filename(source, resolved_base_dir)
    )


def build_lemit_processed_stem(source: str | Path | bytes | BinaryIO, row_count: int) -> str:
    source_name = Path(str(getattr(source, "name", source))).stem
    uf = extract_uf_from_filename(source_name)
    return f"{row_count} refin c6 {uf} (1 ou 0 cartões) sem RL pra disc"


def build_lemit_processed_xlsx_path(
    source: str | Path | bytes | BinaryIO, row_count: int
) -> Path:
    source_path = Path(str(getattr(source, "name", source)))
    return source_path.with_name(f"{build_lemit_processed_stem(source, row_count)}.xlsx")


def build_argus_output_filename_from_stem(
    source: str | Path | bytes | BinaryIO,
    row_count: int,
    base_dir: str | Path | Iterable[str | Path] = ARGUS_DISC_BASE_DIRS,
) -> str:
    processed_stem = build_lemit_processed_stem(source, row_count)
    return f"{build_argus_campaign_name(source, base_dir)} {processed_stem}.csv"


def build_argus_output_path_from_stem(
    source: str | Path | bytes | BinaryIO,
    row_count: int,
    base_dir: str | Path | Iterable[str | Path] = ARGUS_DISC_BASE_DIRS,
    today: date | None = None,
) -> Path:
    source_name = Path(str(getattr(source, "name", source))).stem
    uf = extract_uf_from_filename(source_name)
    resolved_base_dir = resolve_argus_disc_base_dir(base_dir)
    return (
        resolved_base_dir
        / _argus_month_folder_name(today)
        / uf
        / build_argus_output_filename_from_stem(source, row_count, resolved_base_dir)
    )


def extract_uf_from_filename(filename: str) -> str:
    normalized = _normalize_text(filename)
    tokens = [token for token in re.split(r"[^A-Z0-9]+", normalized) if token]
    for idx, token in enumerate(tokens[:-1]):
        if token == "C6" and tokens[idx + 1] in BRAZIL_UFS:
            return tokens[idx + 1]

    ufs = [token for token in tokens if token in BRAZIL_UFS]
    if not ufs:
        raise ValueError("Nao encontrei a UF do estado no nome do arquivo.")
    return ufs[-1]


def find_next_argus_campaign_number(
    uf: str, base_dir: str | Path | Iterable[str | Path] = ARGUS_DISC_BASE_DIRS
) -> int:
    base_path = resolve_argus_disc_base_dir(base_dir)

    pattern = re.compile(rf"\b{re.escape(uf)}\s+(\d+)\b")
    numbers: list[int] = []

    for path in base_path.rglob("*"):
        names = [path.name, *path.parts]
        if not any(_has_uf_token(name, uf) for name in names):
            continue
        for name in names:
            match = pattern.search(_normalize_text(name))
            if match:
                numbers.append(int(match.group(1)))

    return max(numbers, default=0) + 1


def resolve_argus_disc_base_dir(
    base_dir: str | Path | Iterable[str | Path] = ARGUS_DISC_BASE_DIRS,
) -> Path:
    if isinstance(base_dir, (str, Path)):
        candidates = (Path(base_dir),)
    else:
        candidates = tuple(Path(path) for path in base_dir)

    for candidate in candidates:
        if candidate.exists():
            return candidate

    tested_paths = " | ".join(str(path) for path in candidates)
    raise FileNotFoundError(f"Pasta da Argus nao encontrada. Caminhos testados: {tested_paths}")


def find_next_sms_campaign_number(
    base_dir: str | Path = DISPARO_PRO_ENVIADOS_DIR, today: date | None = None
) -> int:
    month_dir = resolve_sms_month_dir(base_dir, today)
    pattern = re.compile(r"\bCAMP\s+(\d+)\b")
    numbers: list[int] = []

    for path in month_dir.glob("*"):
        match = pattern.search(_normalize_text(path.name))
        if match:
            numbers.append(int(match.group(1)))

    return max(numbers, default=0) + 1


def resolve_sms_month_dir(
    base_dir: str | Path = DISPARO_PRO_ENVIADOS_DIR, today: date | None = None
) -> Path:
    base_path = Path(base_dir)
    if not base_path.exists():
        raise FileNotFoundError(f"Pasta do Disparo Pro nao encontrada: {base_path}")

    month_dir = base_path / _argus_month_folder_name(today)
    month_dir.mkdir(parents=True, exist_ok=True)
    return month_dir


def resolve_sms_auto_enviados_dir(
    base_dir: str | Path = DISPARO_PRO_AUTO_ENVIADOS_DIR,
) -> Path:
    auto_dir = Path(base_dir)
    auto_dir.mkdir(parents=True, exist_ok=True)
    return auto_dir


def _load_google_access_token() -> str:
    config_dir = _google_config_dir()
    token_path = config_dir / "token.json"
    token_data = _read_json_file(token_path)
    if token_data and token_data.get("access_token") and token_data.get("expires_at", 0) > time.time() + 60:
        return str(token_data["access_token"])

    if token_data and token_data.get("refresh_token"):
        refreshed = _refresh_google_token(token_data)
        token_path.write_text(json.dumps(refreshed, indent=2), encoding="utf-8")
        return str(refreshed["access_token"])

    client_secret = _load_google_client_secret()
    token_data = _run_google_oauth_flow(client_secret)
    config_dir.mkdir(parents=True, exist_ok=True)
    token_path.write_text(json.dumps(token_data, indent=2), encoding="utf-8")
    return str(token_data["access_token"])


def _load_google_client_secret() -> dict:
    client_secret_path = _find_google_client_secret_path()
    if client_secret_path is None:
        raise GoogleSheetsSetupError(
            "Google Sheets ainda nao configurado. Salvei o XLSX de reserva em auto-enviados. "
            "Para criar .gsheet automaticamente, coloque google_client_secret.json em "
            f"{_google_config_dir()} ou ao lado do executavel."
        )

    data = _read_json_file(client_secret_path)
    if not data:
        raise GoogleSheetsSetupError(f"Credencial Google vazia ou invalida: {client_secret_path}")

    client = data.get("installed") or data.get("web") or data
    if not client.get("client_id") or not client.get("client_secret"):
        raise GoogleSheetsSetupError(f"Credencial Google sem client_id/client_secret: {client_secret_path}")
    return client


def _run_google_oauth_flow(client: dict) -> dict:
    state = secrets.token_urlsafe(24)
    server = HTTPServer(("localhost", 0), _GoogleOAuthCallbackHandler)
    server.auth_code = None
    server.auth_error = None
    server.expected_state = state
    redirect_uri = f"http://localhost:{server.server_port}/"

    auth_uri = client.get("auth_uri", "https://accounts.google.com/o/oauth2/v2/auth")
    query = urllib_parse.urlencode(
        {
            "client_id": client["client_id"],
            "redirect_uri": redirect_uri,
            "response_type": "code",
            "scope": " ".join(GOOGLE_SCOPES),
            "access_type": "offline",
            "prompt": "consent",
            "state": state,
        }
    )
    webbrowser.open(f"{auth_uri}?{query}")
    server.handle_request()

    if server.auth_error:
        raise GoogleSheetsSetupError(f"Login Google cancelado ou recusado: {server.auth_error}")
    if not server.auth_code:
        raise GoogleSheetsSetupError("Nao recebi o codigo de login do Google.")

    token = _request_google_token(
        client,
        {
            "code": server.auth_code,
            "client_id": client["client_id"],
            "client_secret": client["client_secret"],
            "redirect_uri": redirect_uri,
            "grant_type": "authorization_code",
        },
    )
    if "refresh_token" not in token:
        raise GoogleSheetsSetupError(
            "Google nao retornou refresh_token. Apague token.json e tente novamente com prompt de consentimento."
        )
    return token


def _refresh_google_token(token_data: dict) -> dict:
    client = _load_google_client_secret()
    refreshed = _request_google_token(
        client,
        {
            "client_id": client["client_id"],
            "client_secret": client["client_secret"],
            "refresh_token": token_data["refresh_token"],
            "grant_type": "refresh_token",
        },
    )
    refreshed["refresh_token"] = token_data["refresh_token"]
    return refreshed


def _request_google_token(client: dict, payload: dict) -> dict:
    token_uri = client.get("token_uri", "https://oauth2.googleapis.com/token")
    data = urllib_parse.urlencode(payload).encode("utf-8")
    req = urllib_request.Request(
        token_uri,
        data=data,
        headers={"Content-Type": "application/x-www-form-urlencoded"},
        method="POST",
    )
    try:
        with urllib_request.urlopen(req, timeout=120) as response:
            token = json.loads(response.read().decode("utf-8"))
    except urllib_error.HTTPError as exc:
        detail = exc.read().decode("utf-8", errors="replace")
        raise GoogleSheetsSetupError(f"Falha ao autenticar no Google: {detail}") from exc

    token["expires_at"] = time.time() + float(token.get("expires_in", 3600))
    return token


class _GoogleOAuthCallbackHandler(BaseHTTPRequestHandler):
    def do_GET(self) -> None:
        parsed = urllib_parse.urlparse(self.path)
        params = urllib_parse.parse_qs(parsed.query)
        state = params.get("state", [""])[0]
        if state != getattr(self.server, "expected_state", ""):
            self.server.auth_error = "state invalido"
        elif "error" in params:
            self.server.auth_error = params["error"][0]
        else:
            self.server.auth_code = params.get("code", [""])[0]

        self.send_response(200)
        self.send_header("Content-Type", "text/html; charset=utf-8")
        self.end_headers()
        self.wfile.write(
            "<html><body><h2>Login concluido.</h2><p>Pode voltar para o Auto C6BANK.</p></body></html>".encode(
                "utf-8"
            )
        )

    def log_message(self, format: str, *args) -> None:
        return


def _read_json_file(path: Path) -> dict | None:
    if path.exists():
        try:
            return json.loads(path.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            return None
    return None


def _ensure_drive_folder_path(access_token: str, folder_names: tuple[str, ...]) -> str:
    parent_id = "root"
    for folder_name in folder_names:
        folder_id = _find_drive_folder(access_token, folder_name, parent_id)
        if folder_id is None:
            metadata = {
                "name": folder_name,
                "mimeType": "application/vnd.google-apps.folder",
                "parents": [parent_id],
            }
            created = _google_api_request(
                "POST",
                "https://www.googleapis.com/drive/v3/files?fields=id&supportsAllDrives=true",
                access_token,
                metadata,
            )
            folder_id = created["id"]
        parent_id = folder_id
    return parent_id


def _find_drive_folder(access_token: str, name: str, parent_id: str) -> str | None:
    escaped_name = name.replace("\\", "\\\\").replace("'", "\\'")
    escaped_parent = parent_id.replace("\\", "\\\\").replace("'", "\\'")
    query = (
        "mimeType='application/vnd.google-apps.folder' "
        f"and name='{escaped_name}' "
        f"and '{escaped_parent}' in parents "
        "and trashed=false"
    )
    params = urllib_parse.urlencode(
        {
            "q": query,
            "spaces": "drive",
            "fields": "files(id,name)",
            "pageSize": "10",
            "supportsAllDrives": "true",
            "includeItemsFromAllDrives": "true",
        }
    )
    response = _google_api_request(
        "GET",
        f"https://www.googleapis.com/drive/v3/files?{params}",
        access_token,
    )
    files = response.get("files", [])
    return files[0]["id"] if files else None


def _google_api_request(method: str, url: str, access_token: str, payload: dict | None = None) -> dict:
    data = None if payload is None else json.dumps(payload).encode("utf-8")
    headers = {"Authorization": f"Bearer {access_token}"}
    if payload is not None:
        headers["Content-Type"] = "application/json"
    req = urllib_request.Request(url, data=data, headers=headers, method=method)
    try:
        with urllib_request.urlopen(req, timeout=120) as response:
            body = response.read().decode("utf-8")
            return json.loads(body) if body else {}
    except urllib_error.HTTPError as exc:
        detail = exc.read().decode("utf-8", errors="replace")
        raise GoogleSheetsSetupError(f"Falha na API do Google: {detail}") from exc


def _google_config_dir() -> Path:
    return Path.home() / ".auto_c6bank"


def _runtime_dir() -> Path:
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent


def _google_client_secret_candidates() -> tuple[Path, ...]:
    return (
        _runtime_dir() / "google_client_secret.json",
        _runtime_dir() / "credentials.json",
        _google_config_dir() / "google_client_secret.json",
    )


def _find_google_client_secret_path() -> Path | None:
    for path in _google_client_secret_candidates():
        if path.exists():
            return path
    return None


def _google_auto_enviados_folder_id() -> str:
    for path in (
        _runtime_dir() / "google_auto_enviados_folder_id.txt",
        _google_config_dir() / "google_auto_enviados_folder_id.txt",
    ):
        if path.exists():
            folder_id = path.read_text(encoding="utf-8").strip()
            if folder_id:
                return folder_id
    return GOOGLE_DRIVE_AUTO_ENVIADOS_FOLDER_ID


def _argus_month_folder_name(today: date | None = None) -> str:
    current = today or date.today()
    return f"{MONTH_NAMES_PT_BR[current.month]} {current.year}"


def _argus_dataframe_to_csv_bytes(df: pd.DataFrame) -> bytes:
    return df.to_csv(
        index=False,
        sep=";",
        encoding="utf-8-sig",
        quoting=csv.QUOTE_ALL,
    ).encode("utf-8-sig")


def read_processed_c6_file(source: str | Path | bytes | BinaryIO) -> pd.DataFrame:
    source_name = str(getattr(source, "name", ""))
    if isinstance(source, (str, Path)):
        source_name = str(source)

    if source_name.lower().endswith(".xlsx"):
        raw = _read_visible_xlsx(source)
    elif source_name.lower().endswith(".xls"):
        raw = pd.read_excel(source, dtype=str)
    else:
        data = _read_bytes(source)
        last_error: Exception | None = None
        raw = None
        for encoding in ("utf-8-sig", "cp1252", "latin1"):
            try:
                text = data.decode(encoding)
                raw = pd.read_csv(io.StringIO(text), sep=None, engine="python", dtype=str)
                break
            except Exception as exc:
                last_error = exc
        if raw is None:
            raise ValueError(f"Nao consegui ler a planilha para Argus. Ultimo erro: {last_error}")

    raw.columns = [str(column).strip().lstrip("\ufeff") for column in raw.columns]
    columns = {_normalize_header(column): column for column in raw.columns}
    aliases = {
        "CPF": "CPF",
        "TAXA": "TAXA",
        "FEZEM": "FEZ EM",
        "FALTAM": "FALTAM",
        "PARCELA": "PARCELA",
        "TROCO": "TROCO",
        "NOME": "NOME",
        "DATADENASC": "DATA DE NASC",
        "DATANASCIMENTO": "DATA DE NASC",
        "TELEFONEC6": "TELEFONE C6",
        "TELEFONEBASE": "TELEFONE BASE",
    }
    required = (
        "CPF",
        "TAXA",
        "FEZ EM",
        "FALTAM",
        "PARCELA",
        "TROCO",
        "NOME",
        "DATA DE NASC",
        "TELEFONE C6",
        "TELEFONE BASE",
    )

    output = pd.DataFrame()
    for normalized, target in aliases.items():
        if normalized in columns:
            output[target] = raw[columns[normalized]]

    missing = [column for column in required if column not in output.columns]
    if missing:
        raise ValueError(
            "Planilha para Argus sem coluna obrigatoria: " + ", ".join(missing)
        )

    output["CPF"] = output["CPF"].map(_normalize_cpf_key)
    output["TAXA"] = output["TAXA"].map(_parse_br_number)
    output["FEZ EM"] = output["FEZ EM"].map(_parse_br_number).round().astype("Int64")
    output["FALTAM"] = output["FALTAM"].map(_parse_br_number).round().astype("Int64")
    output["PARCELA"] = output["PARCELA"].map(_parse_br_number)
    output["TROCO"] = output["TROCO"].map(_parse_br_number)
    output["NOME"] = output["NOME"].fillna("").astype(str).str.strip()
    output["DATA DE NASC"] = output["DATA DE NASC"].map(_parse_br_date_value)
    output["TELEFONE C6"] = output["TELEFONE C6"].map(_format_identifier_exact)
    output["TELEFONE BASE"] = output["TELEFONE BASE"].map(_format_identifier_exact)
    return output[list(required)]


def _read_visible_xlsx(source: str | Path | bytes | BinaryIO) -> pd.DataFrame:
    from openpyxl import load_workbook

    workbook_source = source
    if not isinstance(source, (str, Path)):
        workbook_source = io.BytesIO(_read_bytes(source))

    workbook = load_workbook(workbook_source, data_only=True, read_only=False)
    worksheet = workbook.active

    visible_rows = []
    for excel_row_number, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
        if worksheet.row_dimensions[excel_row_number].hidden:
            continue
        visible_rows.append(list(row))

    if not visible_rows:
        return pd.DataFrame()

    header_index = None
    for index, row in enumerate(visible_rows):
        if any(value not in (None, "") for value in row):
            header_index = index
            break
    if header_index is None:
        return pd.DataFrame()

    headers = ["" if value is None else str(value).strip() for value in visible_rows[header_index]]
    data_rows = visible_rows[header_index + 1 :]
    return pd.DataFrame(data_rows, columns=headers, dtype=str).dropna(how="all")


def build_filtered_from_csv(
    source: str | Path | bytes | BinaryIO,
    options: FilterOptions,
    phone_base_dir: str | Path | Iterable[str | Path] | None = None,
    lemit_source: str | Path | bytes | BinaryIO | None = None,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    raw = read_c6_csv(source)
    transformed = transform_c6_dataframe(raw)
    filtered = apply_filters(transformed, options)
    if phone_base_dir is not None or lemit_source is not None:
        filtered = enrich_with_phone_sources(
            filtered, phone_base_dir, lemit_source, options.telefone_data_inicial
        )
    filtered = enrich_with_ddb(filtered, source)
    return transformed, filtered


def enrich_with_ddb(
    df: pd.DataFrame,
    source: str | Path | bytes | BinaryIO,
    base_dirs: str | Path | Iterable[str | Path] = DDB_BASE_DIRS,
) -> pd.DataFrame:
    """Preenche DDB cruzando o beneficio da extracao com NB na base da UF."""
    output = df.copy()
    source_name = Path(str(getattr(source, "name", source))).stem
    uf = extract_uf_from_filename(source_name)
    lookup = load_ddb_lookup_for_uf(uf, base_dirs)
    benefit_keys = output["BENEF\u00cdCIO"].map(_digits_only)
    output["BDD"] = benefit_keys.map(lookup).fillna("")
    return output


def load_ddb_lookup_for_uf(
    uf: str,
    base_dirs: str | Path | Iterable[str | Path] = DDB_BASE_DIRS,
) -> dict[str, object]:
    base_dir = resolve_ddb_base_dir(base_dirs)
    candidates = [
        path
        for path in base_dir.rglob("*")
        if path.is_file()
        and path.suffix.lower() in {".csv", ".xlsx", ".xls"}
        and _has_uf_token(path.stem, uf)
    ]
    if not candidates:
        raise FileNotFoundError(f'Nao encontrei a base DDB do estado {uf} em "{base_dir}".')
    if len(candidates) > 1:
        names = ", ".join(sorted(path.name for path in candidates))
        raise ValueError(f"Encontrei mais de uma base DDB para {uf}: {names}")

    raw = _read_ddb_table(candidates[0])
    columns = {_normalize_header(column): column for column in raw.columns}
    benefit_column = next(
        (columns[name] for name in ("NB", "BENEFICIO", "BENEFICIOINSS") if name in columns),
        None,
    )
    ddb_column = columns.get("DDB")
    if benefit_column is None or ddb_column is None:
        raise ValueError(
            f'A base DDB "{candidates[0].name}" precisa conter as colunas NB (ou Beneficio) e DDB.'
        )

    pairs = pd.DataFrame(
        {
            "beneficio": raw[benefit_column].map(_digits_only),
            "DDB": pd.to_datetime(raw[ddb_column], dayfirst=True, errors="coerce").dt.strftime("%d/%m/%Y"),
        }
    ).dropna(subset=["DDB"])
    pairs = pairs.loc[pairs["beneficio"] != ""]
    conflict_counts = pairs.groupby("beneficio", sort=False)["DDB"].nunique()
    conflicts = conflict_counts.loc[conflict_counts > 1].index.tolist()
    if conflicts:
        examples = ", ".join(sorted(conflicts)[:5])
        raise ValueError(
            f'A base DDB de {uf} possui beneficios com DDB conflitante (ex.: {examples}).'
        )
    unique_pairs = pairs.drop_duplicates(subset=["beneficio"], keep="first")
    return dict(zip(unique_pairs["beneficio"], unique_pairs["DDB"]))


def resolve_ddb_base_dir(
    base_dirs: str | Path | Iterable[str | Path] = DDB_BASE_DIRS,
) -> Path:
    candidates = (Path(base_dirs),) if isinstance(base_dirs, (str, Path)) else tuple(map(Path, base_dirs))
    for candidate in candidates:
        if candidate.exists():
            return candidate
    tested = ", ".join(str(path) for path in candidates)
    raise FileNotFoundError(f"Pasta das bases DDB nao encontrada. Caminhos testados: {tested}")


def _read_ddb_table(path: Path) -> pd.DataFrame:
    if path.suffix.lower() in {".xlsx", ".xls"}:
        return pd.read_excel(path, dtype=str)
    data = path.read_bytes()
    last_error: Exception | None = None
    for encoding in ("utf-8-sig", "cp1252", "latin1"):
        try:
            text = data.decode(encoding)
            header = text.splitlines()[0] if text.splitlines() else ""
            separator = max((";", ",", "\t", "|"), key=header.count)
            return pd.read_csv(io.StringIO(text), sep=separator, dtype=str)
        except Exception as exc:
            last_error = exc
    raise ValueError(f"Nao consegui ler a base DDB {path.name}: {last_error}")


def enrich_with_phone_base(df: pd.DataFrame, phone_base_dir: str | Path) -> pd.DataFrame:
    return enrich_with_phone_sources(df, phone_base_dir, None)


def enrich_with_phone_sources(
    df: pd.DataFrame,
    phone_base_dir: str | Path | Iterable[str | Path] | None = None,
    lemit_source: str | Path | bytes | BinaryIO | None = None,
    telefone_data_inicial: date | None = None,
) -> pd.DataFrame:
    output = df.copy()
    if output.empty:
        output["TELEFONE BASE"] = ""
        return output

    lookups: list[dict[str, tuple[str, ...]]] = []
    if phone_base_dir is not None:
        lookups.append(
            load_phone_lookup_for_cpfs(output["CPF"], phone_base_dir, telefone_data_inicial)
        )
    if lemit_source is not None:
        lookups.append(load_lemit_phone_lookup(lemit_source))

    phone_lookup = _merge_phone_lookups(*lookups)
    cpf_keys = output["CPF"].map(_normalize_cpf_key)
    output["TELEFONE BASE"] = cpf_keys.map(
        lambda cpf: " / ".join(phone_lookup.get(cpf, ()))
    ).fillna("")
    return output


def load_phone_lookup_for_cpfs(
    cpfs: Iterable[object],
    phone_base_dir: str | Path | Iterable[str | Path],
    telefone_data_inicial: date | None = None,
) -> dict[str, tuple[str, ...]]:
    base_dirs = resolve_phone_base_dirs(phone_base_dir)

    target_cpfs = {_normalize_cpf_key(cpf) for cpf in cpfs}
    target_cpfs.discard("")
    phone_lists: dict[str, list[str]] = {}
    phone_sets: dict[str, set[str]] = {}

    for base_dir in base_dirs:
        for csv_path in sorted(base_dir.glob("*.csv")):
            for chunk in _read_phone_base_chunks(csv_path, telefone_data_inicial is not None):
                chunk.columns = [str(column).strip() for column in chunk.columns]
                normalized_columns = {_normalize_header(column): column for column in chunk.columns}
                cpf_column = normalized_columns.get("CPF")
                phone_column = normalized_columns.get("TELEFONE")
                date_column = normalized_columns.get("DATADAULTIMAMODIFICACAO")
                if cpf_column is None or phone_column is None:
                    continue

                cpf_series = chunk[cpf_column].map(_normalize_cpf_key)
                phone_series = chunk[phone_column].map(_normalize_phone)
                matched = pd.DataFrame({"cpf": cpf_series, "phone": phone_series})
                if telefone_data_inicial is not None and date_column is not None:
                    date_series = chunk[date_column].map(_parse_br_date_value)
                    min_date = pd.Timestamp(telefone_data_inicial)
                    matched = matched[date_series.isna() | (date_series >= min_date)]
                matched = matched[
                    matched["cpf"].isin(target_cpfs) & (matched["phone"] != "")
                ]

                for cpf, phone in matched.itertuples(index=False):
                    seen = phone_sets.setdefault(cpf, set())
                    if phone in seen:
                        continue
                    seen.add(phone)
                    phone_lists.setdefault(cpf, []).append(phone)

    return {cpf: tuple(phones) for cpf, phones in phone_lists.items()}


def resolve_phone_base_dirs(phone_base_dir: str | Path | Iterable[str | Path]) -> tuple[Path, ...]:
    if isinstance(phone_base_dir, (str, Path)):
        candidates = (Path(phone_base_dir),)
    else:
        candidates = tuple(Path(path) for path in phone_base_dir)

    existing = tuple(candidate for candidate in candidates if candidate.exists())
    if existing:
        return existing

    tested_paths = " | ".join(str(path) for path in candidates)
    raise FileNotFoundError(
        f"Pasta da base de telefones não encontrada. Caminhos testados: {tested_paths}"
    )


def resolve_phone_base_dir(phone_base_dir: str | Path | Iterable[str | Path]) -> Path:
    return resolve_phone_base_dirs(phone_base_dir)[0]

    if isinstance(phone_base_dir, (str, Path)):
        candidates = (Path(phone_base_dir),)
    else:
        candidates = tuple(Path(path) for path in phone_base_dir)

    for candidate in candidates:
        if candidate.exists():
            return candidate

    tested_paths = " | ".join(str(path) for path in candidates)
    raise FileNotFoundError(
        f"Pasta da base de telefones não encontrada. Caminhos testados: {tested_paths}"
    )


def load_lemit_phone_lookup(source: str | Path | bytes | BinaryIO) -> dict[str, tuple[str, ...]]:
    df = read_lemit_file(source)
    normalized_columns = {_normalize_header(column): column for column in df.columns}
    cpf_column = _required_column(normalized_columns, "CPF")
    phone_column = _required_column(normalized_columns, "TELEFONE")
    ddd_column = normalized_columns.get("DDD")
    ranking_column = normalized_columns.get("RANKING")
    score_column = normalized_columns.get("SCORE")

    work_df = df.copy()
    work_df["_cpf"] = work_df[cpf_column].map(_normalize_cpf_key)
    if ddd_column is not None:
        ddd_series = work_df[ddd_column].map(_digits_only)
    else:
        ddd_series = pd.Series([""] * len(work_df), index=work_df.index)
    work_df["_phone"] = [
        _normalize_lemit_phone(ddd, phone)
        for ddd, phone in zip(ddd_series, work_df[phone_column], strict=False)
    ]

    if ranking_column is not None:
        work_df["_ranking"] = pd.to_numeric(work_df[ranking_column], errors="coerce")
    else:
        work_df["_ranking"] = pd.NA
    if score_column is not None:
        work_df["_score"] = pd.to_numeric(work_df[score_column], errors="coerce")
    else:
        work_df["_score"] = pd.NA

    work_df = work_df[(work_df["_cpf"] != "") & (work_df["_phone"] != "")]
    work_df = work_df.sort_values(
        by=["_cpf", "_ranking", "_score"],
        ascending=[True, True, False],
        na_position="last",
        kind="stable",
    )

    phone_lists: dict[str, list[str]] = {}
    phone_sets: dict[str, set[str]] = {}
    for cpf, phone in work_df[["_cpf", "_phone"]].itertuples(index=False):
        seen = phone_sets.setdefault(cpf, set())
        if phone in seen:
            continue
        seen.add(phone)
        phone_lists.setdefault(cpf, []).append(phone)

    return {cpf: tuple(phones) for cpf, phones in phone_lists.items()}


def filter_to_lemit_cpfs(
    df: pd.DataFrame, lemit_source: str | Path | bytes | BinaryIO
) -> pd.DataFrame:
    lemit_df = read_lemit_file(lemit_source)
    normalized_columns = {_normalize_header(column): column for column in lemit_df.columns}
    cpf_column = _required_column(normalized_columns, "CPF")
    lemit_cpfs = {
        cpf
        for cpf in lemit_df[cpf_column].map(_normalize_cpf_key)
        if cpf
    }
    cpf_keys = df["CPF"].map(_normalize_cpf_key)
    return df.loc[cpf_keys.isin(lemit_cpfs)].copy().reset_index(drop=True)


def read_lemit_file(source: str | Path | bytes | BinaryIO) -> pd.DataFrame:
    source_name = str(getattr(source, "name", ""))
    if isinstance(source, (str, Path)):
        source_name = str(source)

    if source_name.lower().endswith(".xlsx"):
        return _read_visible_xlsx(source)
    if source_name.lower().endswith(".xls"):
        return pd.read_excel(source, dtype=str)

    data = _read_bytes(source)
    last_error: Exception | None = None
    for encoding in ("utf-8-sig", "cp1252", "latin1"):
        try:
            text = data.decode(encoding)
            df = pd.read_csv(io.StringIO(text), sep=None, engine="python", dtype=str)
            df.columns = [str(column).strip().lstrip("\ufeff") for column in df.columns]
            return df
        except Exception as exc:
            last_error = exc

    raise ValueError(f"Não consegui ler o retorno Lemiti. Último erro: {last_error}")


def _merge_phone_lookups(
    *lookups: dict[str, tuple[str, ...]]
) -> dict[str, tuple[str, ...]]:
    merged_lists: dict[str, list[str]] = {}
    merged_sets: dict[str, set[str]] = {}
    for lookup in lookups:
        for cpf, phones in lookup.items():
            seen = merged_sets.setdefault(cpf, set())
            current = merged_lists.setdefault(cpf, [])
            for phone in phones:
                if phone in seen:
                    continue
                seen.add(phone)
                current.append(phone)

    return {cpf: tuple(phones) for cpf, phones in merged_lists.items()}


def get_missing_phone_cpfs(df: pd.DataFrame) -> pd.DataFrame:
    if "TELEFONE BASE" not in df.columns:
        missing = df
    else:
        phone_text = df["TELEFONE BASE"].fillna("").astype(str).str.strip()
        missing = df.loc[phone_text == ""]

    output = pd.DataFrame(index=missing.index)
    output["cpf"] = missing["CPF"].map(_normalize_cpf_key)
    return output.loc[output["cpf"] != ""].drop_duplicates(subset=["cpf"])


def missing_phone_cpfs_to_csv_bytes(df: pd.DataFrame) -> bytes:
    return get_missing_phone_cpfs(df).to_csv(index=False, encoding="utf-8-sig").encode(
        "utf-8-sig"
    )


def write_missing_phone_cpfs_csv(df: pd.DataFrame, output_path: str | Path) -> None:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    get_missing_phone_cpfs(df).to_csv(output_path, index=False, encoding="utf-8-sig")


def _read_phone_base_chunks(csv_path: Path, include_modified_date: bool = False):
    last_error: Exception | None = None
    wanted_columns = {"CPF", "TELEFONE"}
    if include_modified_date:
        wanted_columns.add("DATADAULTIMAMODIFICACAO")

    for encoding in ("utf-8-sig", "cp1252", "latin1"):
        try:
            chunks = pd.read_csv(
                csv_path,
                sep=";",
                dtype=str,
                encoding=encoding,
                usecols=lambda column: _normalize_header(column) in wanted_columns,
                chunksize=200_000,
                on_bad_lines="skip",
            )
            for chunk in chunks:
                yield chunk
            return
        except Exception as exc:
            last_error = exc

    raise ValueError(f"Não consegui ler a base de telefones {csv_path}: {last_error}")


def _required_column(columns: dict[str, str], normalized_name: str) -> str:
    if normalized_name not in columns:
        raise ValueError(f"Retorno Lemiti sem coluna obrigatória: {normalized_name}")
    return columns[normalized_name]


def _normalize_header(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value).strip())
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"\s+", "", text).upper()


def _normalize_text(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value).strip())
    text = "".join(char for char in text if not unicodedata.combining(char))
    return text.upper()


def _has_uf_token(value: object, uf: str) -> bool:
    tokens = [token for token in re.split(r"[^A-Z0-9]+", _normalize_text(value)) if token]
    return uf in tokens


def _clean_source_stem(stem: str) -> str:
    return re.sub(r"\s*\(\d+\)\s*$", "", stem).strip()


def _read_bytes(source: str | Path | bytes | BinaryIO) -> bytes:
    if isinstance(source, bytes):
        return source

    if isinstance(source, (str, Path)):
        return Path(source).read_bytes()

    if hasattr(source, "getvalue"):
        return source.getvalue()

    return source.read()


def _digits_only(value: object) -> str:
    if pd.isna(value):
        return ""
    return re.sub(r"\D+", "", str(value).strip())


def _normalize_cpf_key(value: object) -> str:
    digits = _digits_only(value)
    if not digits:
        return ""
    return digits[-11:].zfill(11)


def _normalize_phone(value: object) -> str:
    digits = _digits_only(value)
    if len(digits) < 8:
        return ""
    return digits


def _split_phone_values(value: object) -> tuple[str, ...]:
    phones: list[str] = []
    seen: set[str] = set()
    for part in re.split(r"[/,;|\n]+", _safe_text(value)):
        phone = _normalize_phone(part)
        if not phone or phone in seen:
            continue
        seen.add(phone)
        phones.append(phone[-11:].zfill(11))
    return tuple(phones)


def _normalize_lemit_phone(ddd: object, phone: object) -> str:
    ddd_digits = _digits_only(ddd)
    phone_digits = _digits_only(phone)
    if len(phone_digits) < 8:
        return ""
    if len(phone_digits) <= 9 and ddd_digits:
        return ddd_digits[-2:] + phone_digits
    return phone_digits


def _parse_br_number(value: object) -> float | None:
    if pd.isna(value):
        return None

    text = str(value).strip()
    if not text:
        return None

    text = text.replace("R$", "").replace("\xa0", "").replace(" ", "")
    if "," in text:
        text = text.replace(".", "").replace(",", ".")
    elif text.count(".") > 1:
        parts = text.split(".")
        text = "".join(parts[:-1]) + "." + parts[-1]

    try:
        return float(text)
    except ValueError:
        return None


def _parse_int_series(series: pd.Series) -> pd.Series:
    return series.map(_parse_br_number).round().astype("Int64")


def _safe_text(value: object) -> str:
    if pd.isna(value):
        return ""
    return str(value)


def _google_sheet_cell_value(value: object) -> object:
    if pd.isna(value):
        return ""
    if isinstance(value, (int, float, str, bool)):
        return value
    if isinstance(value, pd.Timestamp):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    return str(value)


def _format_int(value: object) -> str:
    if pd.isna(value):
        return ""
    return str(int(float(value)))


def _format_br_decimal(value: object) -> str:
    if pd.isna(value):
        return ""
    return f"{float(value):.2f}".replace(".", ",")


def _format_br_parcela(value: object) -> str:
    if pd.isna(value):
        return ""

    number = float(value)
    if number.is_integer() and abs(number) < 1000:
        return str(int(number))

    formatted = f"{number:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    if "," in formatted:
        formatted = formatted.rstrip("0").rstrip(",")
    if number.is_integer() and abs(number) >= 1000:
        return formatted + ",00"
    return formatted


def _format_br_accounting(value: object) -> str:
    if pd.isna(value):
        return ""
    number = float(value)
    formatted = f"{abs(number):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    if number < 0:
        return f"R$ ({formatted})"
    return f"R$ {formatted}"


def _format_br_date(value: object) -> str:
    if pd.isna(value):
        return ""
    parsed = _parse_br_date_value(value)
    if pd.isna(parsed):
        return ""
    return pd.to_datetime(parsed).strftime("%d/%m/%Y")


def _parse_br_date_value(value: object):
    if pd.isna(value):
        return pd.NaT

    if isinstance(value, pd.Timestamp):
        return value

    if isinstance(value, datetime):
        return pd.Timestamp(value)

    if isinstance(value, date):
        return pd.Timestamp(value.year, value.month, value.day)

    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            return pd.Timestamp("1899-12-30") + pd.to_timedelta(float(value), unit="D")
        except Exception:
            return pd.NaT

    text = str(value).strip()
    if not text:
        return pd.NaT

    for dayfirst in (True, False):
        parsed = pd.to_datetime(text, dayfirst=dayfirst, errors="coerce")
        if not pd.isna(parsed):
            return parsed

    return pd.NaT


def _format_identifier_exact(value: object) -> str:
    if pd.isna(value):
        return ""

    if isinstance(value, int):
        return str(value)

    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return f"{value:.15g}"

    text = str(value).strip()
    if re.fullmatch(r"\d+\.0", text):
        return text[:-2]
    return text


def _number_or_blank(value: object) -> float | str:
    if pd.isna(value):
        return ""
    return float(value)


def _write_number(worksheet, row: int, col: int, value: object, cell_format) -> None:
    if pd.isna(value):
        worksheet.write_blank(row, col, None, cell_format)
    else:
        worksheet.write_number(row, col, float(value), cell_format)


def _write_digits_number(worksheet, row: int, col: int, value: object, cell_format) -> None:
    digits = _digits_only(value)
    if not digits:
        worksheet.write_blank(row, col, None, cell_format)
    else:
        worksheet.write_number(row, col, int(digits), cell_format)


def _write_date(worksheet, row: int, col: int, value: object, cell_format) -> None:
    if pd.isna(value):
        worksheet.write_blank(row, col, None, cell_format)
        return

    if isinstance(value, pd.Timestamp):
        value = value.to_pydatetime()
    elif isinstance(value, date) and not isinstance(value, datetime):
        value = datetime(value.year, value.month, value.day)

    worksheet.write_datetime(row, col, value, cell_format)


def _parse_cli_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Filtra exportação C6 e gera XLSX formatado.")
    parser.add_argument("--input", required=True, help="Caminho da Extração DTF CSV.")
    parser.add_argument("--output", required=True, help="Caminho do XLSX final.")
    parser.add_argument(
        "--phone-base-dir",
        default="",
        help=(
            "Pasta dos CSVs da base de telefones. Se ficar vazio, tenta G: e depois H:."
        ),
    )
    parser.add_argument("--lemit-input", default="", help="Caminho do retorno Lemiti.")
    parser.add_argument("--faltam", default="", help="Prazo(s) FALTAM. Ex.: 67 ou 67,72.")
    parser.add_argument(
        "--faltam-modo",
        default="maior",
        choices=["maior", "a_partir", "ate", "exato"],
        help="Como aplicar o filtro FALTAM.",
    )
    parser.add_argument("--taxa-min", type=float, default=0.0, help="Taxa mínima.")
    parser.add_argument("--troco-min", type=float, default=0.0, help="Troco mínimo, filtro maior que.")
    parser.add_argument("--data-inicial", default="01/01/1954", help="Data inicial.")
    parser.add_argument("--data-final", default="31/12/1995", help="Data final.")
    return parser.parse_args()


def main() -> None:
    args = _parse_cli_args()
    options = FilterOptions(
        faltam=parse_faltam_values(args.faltam),
        faltam_modo=args.faltam_modo,
        taxa_min=args.taxa_min,
        troco_min=args.troco_min,
        data_inicial=parse_date_value(args.data_inicial),
        data_final=parse_date_value(args.data_final),
    )
    lemit_input = args.lemit_input or None
    _, filtered = build_filtered_from_csv(
        args.input,
        options,
        args.phone_base_dir or DEFAULT_PHONE_BASE_DIRS,
        lemit_input,
    )
    write_xlsx(filtered, args.output)
    print(f"Linhas exportadas: {len(filtered)}")
    print(f"Arquivo: {Path(args.output).resolve()}")


if __name__ == "__main__":
    main()
